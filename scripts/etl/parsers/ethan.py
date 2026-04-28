"""Parser fournisseur ETHAN — factures Excel (.xlsx).

Format source : xlsx consolidé avec 2 feuilles :
  - En-têtes Factures : Numero_Facture | Date | Client | Code_Client | Adresse_Client
                       | Montant_HT | TVA | Total_TTC | Total_Verse | Reste_Du
  - Lignes Factures  : Numero_Facture | Date | Client | CODE | Reference | Qte
                       | Prix_HT | Prix_TTC | Remise_Pct | TVA | Tot_HT | Tot_TTC

Contrat public (identique METRO/TAIYAT — ADR-08) :
  - parse(fichier) → list[LigneParsee]
  - parse_facture(fichier) → tuple[list[LigneParsee], FactureMetadata]
  - parse_with_facture_metrics(fichier) → tuple[list[LigneParsee], dict]

Un xlsx peut contenir N factures. Pour respecter le contrat 1-file=1-facture,
les fonctions acceptent un `invoice_number` optionnel. À défaut, la PREMIÈRE
facture du fichier est retournée. Pour batch import : utiliser list_invoices()
puis itérer.

Pas d'EAN dans le format ETHAN (CODE est un ID interne vendor).
Le `client` ("L'INCONTOURNABLE" ou "Sté NOUTAM") détermine `target_tenant_id`.
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from datetime import date as date_cls, datetime
from pathlib import Path
from typing import Optional

from app.etl_types import FactureMetadata, LigneParsee

logger = logging.getLogger(__name__)

_SOURCE_FOURNISSEUR = "ETHAN"

# Routing multi-tenant — INCONTOURNABLE (restaurant) / NOUTAM (épicerie)
_CLIENT_TO_TENANT: dict[str, int] = {
    "INCONTOURNABLE": 3,
    "NOUTAM": 2,
}

_SHEET_HEADERS = "En-têtes Factures"
_SHEET_LINES = "Lignes Factures"


@dataclass
class _InvoiceHeader:
    numero: str
    date: Optional[date_cls]
    client: str
    client_code: Optional[str]
    adresse: Optional[str]
    montant_ht_cts: Optional[int]
    montant_tva_cts: Optional[int]
    montant_ttc_cts: Optional[int]


def _require_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError as exc:
        raise ImportError(
            "openpyxl requis pour parser ETHAN : pip install openpyxl"
        ) from exc


def _euros_to_cts(val) -> Optional[int]:
    if val is None:
        return None
    try:
        return round(float(val) * 100)
    except (ValueError, TypeError):
        return None


def _parse_date(raw) -> Optional[date_cls]:
    if raw is None:
        return None
    if isinstance(raw, date_cls):
        return raw if not isinstance(raw, datetime) else raw.date()
    if isinstance(raw, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(raw, fmt).date()
            except ValueError:
                continue
    return None


def _normalize_client(client: Optional[str]) -> Optional[str]:
    """Mappe 'Sté NOUTAM' / "L'INCONTOURNABLE" → clé normalisée NOUTAM|INCONTOURNABLE."""
    if not client:
        return None
    up = client.upper()
    if "INCONTOURNABLE" in up:
        return "INCONTOURNABLE"
    if "NOUTAM" in up:
        return "NOUTAM"
    return None


def _load_headers(ws) -> dict[str, _InvoiceHeader]:
    """Lit la feuille 'En-têtes Factures' → {numero_facture: _InvoiceHeader}."""
    headers: dict[str, _InvoiceHeader] = {}
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return headers
    col_index = {name: i for i, name in enumerate(rows[0]) if name}

    def get(row, name):
        idx = col_index.get(name)
        return row[idx] if idx is not None and idx < len(row) else None

    for row in rows[1:]:
        num = get(row, "Numero_Facture")
        if not num:
            continue
        num = str(num).strip()
        ht = _euros_to_cts(get(row, "Montant_HT"))
        tva = _euros_to_cts(get(row, "TVA"))
        ttc = _euros_to_cts(get(row, "Total_TTC"))
        headers[num] = _InvoiceHeader(
            numero=num,
            date=_parse_date(get(row, "Date")),
            client=str(get(row, "Client") or ""),
            client_code=str(get(row, "Code_Client")) if get(row, "Code_Client") else None,
            adresse=str(get(row, "Adresse_Client")) if get(row, "Adresse_Client") else None,
            montant_ht_cts=ht,
            montant_tva_cts=tva,
            montant_ttc_cts=ttc,
        )
    return headers


def _load_lines_for_invoice(ws, numero_facture: str) -> list[LigneParsee]:
    """Lit les lignes de la feuille 'Lignes Factures' filtrées par numero_facture."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    col_index = {name: i for i, name in enumerate(rows[0]) if name}

    def get(row, name):
        idx = col_index.get(name)
        return row[idx] if idx is not None and idx < len(row) else None

    lignes: list[LigneParsee] = []
    for row in rows[1:]:
        num = get(row, "Numero_Facture")
        if not num or str(num).strip() != numero_facture:
            continue

        ref = get(row, "Reference")
        if not ref:
            continue

        qte_raw = get(row, "Qte")
        try:
            qte = float(qte_raw) if qte_raw is not None else None
        except (ValueError, TypeError):
            qte = None

        prix_ttc = _euros_to_cts(get(row, "Prix_TTC"))
        tot_ht = _euros_to_cts(get(row, "Tot_HT"))
        tot_ttc = _euros_to_cts(get(row, "Tot_TTC"))

        tva_pct = get(row, "TVA")
        try:
            tva_centieme = round(float(tva_pct) * 100) if tva_pct is not None else None
        except (ValueError, TypeError):
            tva_centieme = None

        code = get(row, "CODE")
        article_fournisseur = str(code) if code is not None else None

        lignes.append(LigneParsee(
            designation=str(ref).strip(),
            unite_base="U",
            source_fournisseur=_SOURCE_FOURNISSEUR,
            ean=None,
            quantite=qte,
            prix_unitaire_cts=prix_ttc,
            montant_ht_cts=tot_ht,
            montant_ttc_cts=tot_ttc,
            taux_tva_centieme=tva_centieme,
            article_fournisseur=article_fournisseur,
            designation_raw=str(ref).strip(),
        ))
    return lignes


def list_invoices(fichier: str) -> list[str]:
    """Retourne la liste ordonnée des numeros de facture présents dans le xlsx."""
    openpyxl = _require_openpyxl()
    path = Path(fichier)
    if not path.exists():
        raise FileNotFoundError(f"Fichier introuvable: {fichier}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if _SHEET_HEADERS not in wb.sheetnames:
        return []
    headers = _load_headers(wb[_SHEET_HEADERS])
    return sorted(headers.keys(), key=lambda x: headers[x].date or date_cls.min)


def parse_with_facture_metrics(
    fichier: str, invoice_number: Optional[str] = None,
) -> tuple[list[LigneParsee], dict]:
    """Parse un xlsx ETHAN et retourne lignes + métriques pour une facture.

    Si invoice_number est None, la première facture disponible est utilisée.
    """
    openpyxl = _require_openpyxl()
    path = Path(fichier)
    if not path.exists():
        raise FileNotFoundError(f"Fichier introuvable: {fichier}")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if _SHEET_HEADERS not in wb.sheetnames or _SHEET_LINES not in wb.sheetnames:
        raise ValueError(
            f"Feuilles attendues absentes. Trouvées : {wb.sheetnames}"
        )

    headers = _load_headers(wb[_SHEET_HEADERS])
    if not headers:
        return [], {"file": str(path), "error": "no_headers"}

    if invoice_number is None:
        invoice_number = next(iter(headers))
    elif invoice_number not in headers:
        raise ValueError(
            f"Facture {invoice_number!r} introuvable. Disponibles : {list(headers.keys())[:5]}..."
        )

    header = headers[invoice_number]
    lignes = _load_lines_for_invoice(wb[_SHEET_LINES], invoice_number)

    # Cohérence totaux
    computed_ttc_cts = sum((l.montant_ttc_cts or 0) for l in lignes)
    declared_ttc_cts = header.montant_ttc_cts or 0
    ecart_cts = declared_ttc_cts - computed_ttc_cts
    ecart_eur = ecart_cts / 100.0

    line_outliers = []
    for l in lignes:
        if l.quantite is None or l.prix_unitaire_cts is None:
            continue
        expected_cts = round(l.quantite * l.prix_unitaire_cts)
        delta = abs(expected_cts - (l.montant_ttc_cts or 0))
        if delta > 2:  # >2 centimes → outlier
            line_outliers.append({
                "designation": l.designation,
                "expected_cts": expected_cts,
                "actual_cts": l.montant_ttc_cts,
                "delta_cts": delta,
            })

    client_name = _normalize_client(header.client)

    quality = 0
    if header.numero:
        quality += 20
    if header.date:
        quality += 20
    if declared_ttc_cts:
        quality += 20
    if lignes:
        quality += 20 * (1 - len(line_outliers) / max(len(lignes), 1))
    if abs(ecart_cts) <= 5:  # 5 centimes
        quality += 20

    metrics = {
        "file": str(path),
        "invoice_number": header.numero,
        "invoice_date": header.date.isoformat() if header.date else None,
        "client_name": client_name,
        "client_code": header.client_code,
        "line_count_output": len(lignes),
        "line_count_rejected": 0,
        "line_coherence_outlier_count": len(line_outliers),
        "line_coherence_outliers": line_outliers,
        "total_ttc_declared": declared_ttc_cts / 100.0 if declared_ttc_cts else None,
        "total_ttc_computed": computed_ttc_cts / 100.0,
        "total_ht_declared_cts": header.montant_ht_cts,
        "ecart_ttc": ecart_eur,
        "ecart_ttc_abs": abs(ecart_eur),
        "is_total_coherent": abs(ecart_cts) <= 5,
        "quality_score": round(quality, 2),
    }

    return lignes, metrics


def parse_facture(
    fichier: str, invoice_number: Optional[str] = None,
) -> tuple[list[LigneParsee], FactureMetadata]:
    """Parse une facture ETHAN et retourne lignes + FactureMetadata (ADR-25)."""
    lignes, metrics = parse_with_facture_metrics(fichier, invoice_number)

    openpyxl = _require_openpyxl()
    wb = openpyxl.load_workbook(fichier, read_only=True, data_only=True)
    headers = _load_headers(wb[_SHEET_HEADERS])
    if invoice_number is None:
        invoice_number = next(iter(headers))
    header = headers.get(invoice_number)

    client_name = metrics.get("client_name")
    target_tenant_id = _CLIENT_TO_TENANT.get(client_name) if client_name else None

    metadata = FactureMetadata(
        vendor_code=_SOURCE_FOURNISSEUR,
        numero_facture=header.numero if header else None,
        date_facture=header.date if header else None,
        montant_ht_total=header.montant_ht_cts if header else None,
        montant_tva_total=header.montant_tva_cts if header else None,
        montant_ttc_total=header.montant_ttc_cts if header else None,
        quality_score=min(round(metrics.get("quality_score", 0)), 100),
        ecart_reconciliation=metrics.get("ecart_ttc"),
        lignes_brutes=len(lignes),
        client_name=client_name,
        target_tenant_id=target_tenant_id,
    )
    return lignes, metadata


def parse(fichier: str, invoice_number: Optional[str] = None) -> list[LigneParsee]:
    """Parse simple — retourne la liste des lignes d'une facture ETHAN."""
    lignes, _ = parse_with_facture_metrics(fichier, invoice_number)
    return lignes


__all__ = [
    "parse", "parse_facture", "parse_with_facture_metrics",
    "list_invoices",
]
